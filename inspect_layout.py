from pptx import Presentation
import sys

def inspect_slide(file_path):
    prs = Presentation(file_path)
    slide = prs.slides[3]
    for j, shape in enumerate(slide.shapes):
        if hasattr(shape, "text_frame") and shape.text_frame:
            text = shape.text.replace('\n', ' ')
            print(f"Shape {j}: x={shape.left}, y={shape.top}, w={shape.width}, h={shape.height} | Text: {text[:30]}")

if __name__ == "__main__":
    inspect_slide(sys.argv[1])
