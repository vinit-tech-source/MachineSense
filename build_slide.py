from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
import sys

def set_text(shape, text, font_name="Calibri", font_size=14, bold=False):
    shape.text = text
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = font_name
            run.font.size = Pt(font_size)
            run.font.bold = bold

def run_build(file_path, out_path):
    prs = Presentation(file_path)
    slide = prs.slides[3]
    
    # 1. Main statement (Shape 5)
    main_text = "An adaptive machine-monitoring and AI system that turns available factory data into contextual insights, anomaly detection, and actionable recommendations."
    set_text(slide.shapes[5], main_text, font_size=18)

    # Re-use shapes for columns to retain theme
    col1 = slide.shapes[7]
    col2 = slide.shapes[12]
    col3 = slide.shapes[16]
    col4 = slide.shapes[21]
    kd_shape = slide.shapes[11]
    bot_shape = slide.shapes[26]

    # Map positions
    y_cols = 2700000
    h_cols = 2300000
    w_cols = 2600000
    
    col1.left = 457200
    col1.top = y_cols
    col1.width = w_cols
    col1.height = h_cols
    
    col2.left = 3200000
    col2.top = y_cols
    col2.width = w_cols
    col2.height = h_cols
    
    col3.left = 5942800
    col3.top = y_cols
    col3.width = w_cols
    col3.height = h_cols
    
    col4.left = 8685600
    col4.top = y_cols
    col4.width = w_cols
    col4.height = h_cols
    
    # Set Texts with formatting
    c1_txt = "01 — CONNECT\nCollect machine data\n• Voltage, current, power, energy\n• Production / part count\n• Reject count\n• Temperature, vibration\n• Optional RPM and pressure"
    set_text(col1, c1_txt, font_size=14)
    
    c2_txt = "02 — UNDERSTAND\nLearn machine behaviour\n• Learn normal behaviour for each machine\n• Identify operating state\n• Establish machine-specific baseline\n• Calculate energy per good unit (SEC)"
    set_text(col2, c2_txt, font_size=14)
    
    c3_txt = "03 — DETECT\nIdentify meaningful anomalies\n• ML-based anomaly detection\n• Compare live data with baseline\n• Correlate multiple signals\n• Consider production and operating conditions\n• Reduce false alarms"
    set_text(col3, c3_txt, font_size=14)
    
    c4_txt = "04 — ACT\nTurn insights into action\n• Identify likely cause\n• Show confidence and severity\n• Provide supporting evidence\n• Recommend safe operator action\n• Estimate energy and cost impact"
    set_text(col4, c4_txt, font_size=14)

    # Key Differentiator
    kd_shape.left = 457200
    kd_shape.top = 5100000
    kd_shape.width = 10800000
    kd_shape.height = 700000
    kd_text = "Key Differentiator:\nMachine-specific contextual intelligence. The system does not rely on fixed universal thresholds. It learns each machine's normal behaviour and combines baseline analysis, SEC, anomaly scores, and supporting signals before generating a recommendation."
    set_text(kd_shape, kd_text, font_size=14)

    # Bottom Statement
    bot_shape.left = 457200
    bot_shape.top = 5900000
    bot_shape.width = 10800000
    bot_shape.height = 500000
    bot_text = "Measure → Understand → Detect → Explain → Act"
    set_text(bot_shape, bot_text, font_size=16, bold=True)
    # Center text
    for paragraph in bot_shape.text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER

    # Delete all other shapes in the body area (from shape index 6 to 31, except the ones we kept)
    keep_indices = [5, 7, 11, 12, 16, 21, 26]
    # To delete shapes safely without messing up indices, find the actual shape objects
    shapes_to_delete = []
    for i in range(6, 32):
        if i not in keep_indices and i != 32: # 32 is footer
            if i < len(slide.shapes):
                shapes_to_delete.append(slide.shapes[i])
                
    for shape in shapes_to_delete:
        element = shape.element
        element.getparent().remove(element)

    prs.save(out_path)

if __name__ == "__main__":
    run_build(sys.argv[1], sys.argv[2])
