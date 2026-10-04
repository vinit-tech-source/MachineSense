from pptx import Presentation

def inspect_fonts(file_path):
    prs = Presentation(file_path)
    slide1 = prs.slides[0]
    print("--- Slide 1 Fonts ---")
    for shape in slide1.shapes:
        if hasattr(shape, "text_frame") and shape.text_frame:
            print(f"Shape text snippet: '{shape.text[:30].replace(chr(10), ' ')}'")
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    font = r.font
                    print(f"  Run text: '{r.text}' | Font: {font.name}, Size: {font.size.pt if font.size else 'None'}")

if __name__ == "__main__":
    import sys
    inspect_fonts(sys.argv[1])
