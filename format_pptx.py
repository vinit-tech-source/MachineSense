from pptx import Presentation
from pptx.util import Pt

def format_presentation(file_path, out_path):
    prs = Presentation(file_path)
    
    corner_phrases = ["YUVA YODHA", "Energy Tech Hackathon 2026", "SCHNEIDER ELECTRIC", "CHALLENGE 04"]
    footer_phrases = ["Team MachineSense", "01 / 12", "02 / 12", "03 / 12", "04 / 12", "05 / 12", "06 / 12", "07 / 12", "08 / 12", "09 / 12", "10 / 12", "11 / 12", "12 / 12"]

    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text_frame") and shape.text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        # 1. Better, modern font name for the entire presentation
                        run.font.name = "Segoe UI"
                        
                        text = run.text.strip()
                        if not text:
                            continue

                        # 2. Fix the upper corners (size 12)
                        is_corner = any(phrase in text for phrase in corner_phrases)
                        is_footer = any(phrase in text for phrase in footer_phrases)

                        if is_corner:
                            run.font.size = Pt(12)
                            if "YUVA" in text or "SCHNEIDER" in text:
                                run.font.bold = True
                        elif is_footer:
                            run.font.size = Pt(10)
                        else:
                            # 3. Better sizes for the rest of the PPT
                            # If size is explicitly set and it's too small for a presentation, bump it
                            if run.font.size is not None:
                                current_pt = run.font.size.pt
                                if current_pt > 30:
                                    # It's a title, ensure it's punchy
                                    run.font.size = Pt(40)
                                    run.font.bold = True
                                elif current_pt < 16:
                                    # It's body text that is too small, make it readable
                                    run.font.size = Pt(18)
                            else:
                                # Inherited size, we can leave it or force a standard body size
                                # For safety, if it's inherited, it's usually defined by the slide master
                                pass

                        # Specific check for the team members we just added to make sure they look good
                        if "Vinit Shinde" in text or "Harsh Vyavahare" in text:
                            run.font.size = Pt(22)
                            run.font.bold = True

    prs.save(out_path)
    print("Formatted PPTX saved to", out_path)

if __name__ == "__main__":
    import sys
    format_presentation(sys.argv[1], sys.argv[2])
