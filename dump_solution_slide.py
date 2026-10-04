from pptx import Presentation
import sys

def dump_slide(file_path):
    prs = Presentation(file_path)
    for i, slide in enumerate(prs.slides):
        is_solution = False
        for shape in slide.shapes:
            if hasattr(shape, "text_frame") and shape.text_frame:
                if "Our Solution" in shape.text:
                    is_solution = True
        if is_solution:
            print(f"--- Slide {i+1} ---")
            for j, shape in enumerate(slide.shapes):
                if hasattr(shape, "text_frame") and shape.text_frame:
                    print(f"Shape {j}:")
                    print(shape.text)
                    print("-" * 20)

if __name__ == "__main__":
    dump_slide(sys.argv[1])
