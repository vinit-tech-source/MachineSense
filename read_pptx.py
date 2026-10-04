from pptx import Presentation
import sys

def dump_pptx_text(file_path):
    try:
        prs = Presentation(file_path)
        for i, slide in enumerate(prs.slides):
            print(f"--- Slide {i+1} ---")
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    print(shape.text.encode('utf-8', 'replace').decode('utf-8'))
    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        dump_pptx_text(sys.argv[1])
    else:
        print("Please provide pptx file path")
